"""CSV/XLSX product-catalog import (preview dry-run + transactional commit)."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, BinaryIO

from .catalog import ProductCatalog, normalize_code, normalize_identity_part

# Conservative documented import limits (server-side enforcement).
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MiB
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_COLUMNS = 64
MAX_IMPORT_SHEETS = 20
MAX_ZIP_MEMBERS = 200
MAX_ZIP_UNCOMPRESSED_TOTAL = 25 * 1024 * 1024  # 25 MiB
MAX_ZIP_MEMBER_UNCOMPRESSED = 10 * 1024 * 1024  # 10 MiB
MAX_ZIP_COMPRESSION_RATIO = 100.0
FORMULA_PREFIXES = ("=", "+", "-", "@")

CANONICAL_FIELDS = (
    "product_name",
    "manufacturer",
    "catalog_number",
    "cas_number",
    "primary_code",
    "alternate_codes",
    "category_id",
)

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "product_name": (
        "product",
        "product name",
        "name",
        "chemical",
        "substance",
        "reagent",
    ),
    "manufacturer": (
        "manufacturer",
        "supplier",
        "vendor",
        "producer",
        "hersteller",
        "lieferant",
    ),
    "catalog_number": (
        "catalogue no",
        "catalogue number",
        "catalog no",
        "catalog number",
        "cat no",
        "cat. no.",
        "cat. no",
        "article number",
        "artikelnummer",
        "bestellnummer",
    ),
    "cas_number": (
        "cas",
        "cas no",
        "cas number",
        "cas-nummer",
        "cas nummer",
    ),
    "primary_code": (
        "code",
        "barcode",
        "qr code",
        "qr",
        "ean",
        "gtin",
        "scanned code",
        "product code",
    ),
    "alternate_codes": (
        "alternate codes",
        "alt codes",
        "alternate barcodes",
        "secondary codes",
    ),
    "category_id": (
        "category",
        "resource category",
        "type",
        "kategorie",
        "category_id",
        "category id",
    ),
}


def neutralize_csv_cell(value: Any) -> str:
    """Neutralize spreadsheet formula injection while preserving data text.

    Cells beginning with =, +, -, or @ are prefixed with a single quote so
    spreadsheet apps treat them as literal text. Leading zeroes in codes are
    otherwise preserved (they do not start with formula prefixes).
    """
    text_value = "" if value is None else str(value)
    if text_value.startswith(FORMULA_PREFIXES):
        return "'" + text_value
    return text_value


def _norm_header(value: str) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("_", " ").replace("-", " ")
    text = re.sub(r"\s+", " ", text)
    text = text.rstrip(".")
    return text


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    """Map canonical fields → source header names (or None)."""
    normalized = {_norm_header(h): h for h in headers if str(h or "").strip()}
    mapping: dict[str, str | None] = {name: None for name in CANONICAL_FIELDS}
    used: set[str] = set()
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            key = _norm_header(alias)
            if key in normalized and normalized[key] not in used:
                mapping[field_name] = normalized[key]
                used.add(normalized[key])
                break
    return mapping


@dataclass
class RowIssue:
    row: int
    field: str
    reason: str

    def to_dict(self) -> dict[str, Any]:
        return {"row": self.row, "field": self.field, "reason": self.reason}


@dataclass
class PreparedRow:
    row_number: int
    values: dict[str, Any]
    action: str  # new | update | duplicate | skip | conflict | error
    existing_id: int | None = None
    issues: list[RowIssue] = field(default_factory=list)


@dataclass
class ImportPreview:
    filename: str
    format: str
    sheets: list[str]
    selected_sheet: str | None
    headers: list[str]
    suggested_mapping: dict[str, str | None]
    mapping: dict[str, str | None]
    total_rows: int
    blank_rows: int
    valid_rows: int
    new_products: int
    updates: int
    exact_duplicates: int
    conflicts: int
    skipped: int
    errors: list[RowIssue]
    rows: list[PreparedRow]

    def summary(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "format": self.format,
            "sheets": self.sheets,
            "selected_sheet": self.selected_sheet,
            "headers": self.headers,
            "suggested_mapping": self.suggested_mapping,
            "mapping": self.mapping,
            "total_rows": self.total_rows,
            "blank_rows": self.blank_rows,
            "valid_rows": self.valid_rows,
            "new_products": self.new_products,
            "updates": self.updates,
            "exact_duplicates": self.exact_duplicates,
            "conflicts": self.conflicts,
            "skipped": self.skipped,
            "error_count": len(self.errors),
            "errors": [e.to_dict() for e in self.errors[:200]],
            "sample_rows": [
                {
                    "row": r.row_number,
                    "action": r.action,
                    "values": r.values,
                    "existing_id": r.existing_id,
                    "issues": [i.to_dict() for i in r.issues],
                }
                for r in self.rows[:50]
            ],
            "limits": {
                "max_upload_bytes": MAX_UPLOAD_BYTES,
                "max_rows": MAX_IMPORT_ROWS,
                "max_columns": MAX_IMPORT_COLUMNS,
                "max_sheets": MAX_IMPORT_SHEETS,
            },
        }

    def error_report_csv(self) -> str:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["row", "field", "reason"])
        for err in self.errors:
            writer.writerow(
                [
                    neutralize_csv_cell(err.row),
                    neutralize_csv_cell(err.field),
                    neutralize_csv_cell(err.reason),
                ]
            )
        for row in self.rows:
            for issue in row.issues:
                if issue not in self.errors:
                    writer.writerow(
                        [
                            neutralize_csv_cell(issue.row),
                            neutralize_csv_cell(issue.field),
                            neutralize_csv_cell(issue.reason),
                        ]
                    )
        return buf.getvalue()


def validate_zip_workbook(data: bytes) -> None:
    """Inspect XLSX ZIP members before openpyxl; reject bombs and unsafe paths."""
    if len(data) < 4 or data[:2] != b"PK":
        raise ValueError("File is not a valid XLSX workbook (missing ZIP signature)")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError("File is not a valid XLSX workbook (malformed ZIP)") from exc

    with zf:
        names = zf.namelist()
        if len(names) > MAX_ZIP_MEMBERS:
            raise ValueError(f"XLSX has too many ZIP members (max {MAX_ZIP_MEMBERS})")

        total_uncompressed = 0
        has_content_types = False
        has_workbook = False
        for info in zf.infolist():
            name = info.filename.replace("\\", "/")
            if name.startswith("/") or name.startswith("\\") or ".." in name.split("/"):
                raise ValueError("XLSX contains unsafe path entries")
            if info.flag_bits & 0x1:
                raise ValueError("Encrypted XLSX archives are not supported")
            if name.endswith("/"):
                continue
            uncompressed = int(info.file_size)
            compressed = max(int(info.compress_size), 1)
            if uncompressed > MAX_ZIP_MEMBER_UNCOMPRESSED:
                raise ValueError("XLSX member exceeds uncompressed size limit")
            ratio = uncompressed / compressed
            if ratio > MAX_ZIP_COMPRESSION_RATIO and uncompressed > 1024 * 1024:
                raise ValueError("XLSX compression ratio exceeds safety limit")
            total_uncompressed += uncompressed
            if total_uncompressed > MAX_ZIP_UNCOMPRESSED_TOTAL:
                raise ValueError("XLSX total uncompressed size exceeds safety limit")
            lower = name.lower()
            if lower == "[content_types].xml":
                has_content_types = True
            if lower.startswith("xl/") and "workbook" in lower:
                has_workbook = True

        if not has_content_types or not has_workbook:
            raise ValueError("File is not a valid OOXML workbook")


def detect_upload_format(filename: str, data: bytes) -> str:
    """Validate extension against content signature; return csv|xlsx."""
    lower = (filename or "").lower().strip()
    if lower.endswith(".xls") and not lower.endswith(".xlsx"):
        raise ValueError("Legacy .xls is not supported — save as .xlsx or .csv")

    is_zip = len(data) >= 4 and data[:2] == b"PK"
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("Legacy .xls is not supported — save as .xlsx or .csv")

    if lower.endswith(".xlsx"):
        if not is_zip:
            raise ValueError("Extension/signature mismatch: .xlsx is not a ZIP workbook")
        validate_zip_workbook(data)
        return "xlsx"

    if lower.endswith(".csv"):
        if is_zip:
            raise ValueError("Extension/signature mismatch: .csv looks like a ZIP archive")
        if b"\x00" in data[:8192]:
            raise ValueError("CSV contains binary/NUL data")
        return "csv"

    if is_zip:
        validate_zip_workbook(data)
        return "xlsx"
    if b"\x00" not in data[:8192]:
        return "csv"
    raise ValueError("Unsupported file type — upload .csv or .xlsx")


def list_xlsx_sheets(data: bytes) -> list[str]:
    validate_zip_workbook(data)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = list(wb.sheetnames)
        if len(sheets) > MAX_IMPORT_SHEETS:
            raise ValueError(f"Too many worksheets (max {MAX_IMPORT_SHEETS})")
        return sheets
    finally:
        wb.close()


def _read_csv_rows(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Unable to decode CSV as UTF-8 or Latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError("CSV has no header row")
    headers = [str(h) for h in reader.fieldnames if h is not None]
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ValueError(f"Too many columns (max {MAX_IMPORT_COLUMNS})")
    rows: list[dict[str, str]] = []
    for raw in reader:
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValueError(f"Too many data rows (max {MAX_IMPORT_ROWS})")
        rows.append({str(k): "" if v is None else str(v) for k, v in raw.items() if k is not None})
    return headers, rows


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Preserve barcode-like integers without trailing .0, keep as int string
        return str(int(value))
    return str(value).strip()


def _read_xlsx_rows(data: bytes, sheet_name: str | None) -> tuple[list[str], list[dict[str, str]], list[str]]:
    validate_zip_workbook(data)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = list(wb.sheetnames)
        if not sheets:
            raise ValueError("XLSX has no sheets")
        if len(sheets) > MAX_IMPORT_SHEETS:
            raise ValueError(f"Too many worksheets (max {MAX_IMPORT_SHEETS})")
        name = sheet_name or sheets[0]
        if name not in sheets:
            raise ValueError(f"Sheet not found: {name}")
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration as exc:
            raise ValueError("Selected sheet is empty") from exc
        if header_row is None:
            raise ValueError("Selected sheet is empty")
        if len(header_row) > MAX_IMPORT_COLUMNS:
            raise ValueError(f"Too many columns (max {MAX_IMPORT_COLUMNS})")
        headers = [_cell_str(h) or f"column_{idx+1}" for idx, h in enumerate(header_row)]
        # Deduplicate empty/duplicate headers
        seen: dict[str, int] = {}
        unique_headers: list[str] = []
        for h in headers:
            base = h or "column"
            if base in seen:
                seen[base] += 1
                unique_headers.append(f"{base}_{seen[base]}")
            else:
                seen[base] = 0
                unique_headers.append(base)
        rows: list[dict[str, str]] = []
        for values in it:
            if values is None:
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ValueError(f"Too many data rows (max {MAX_IMPORT_ROWS})")
            row = {
                unique_headers[i]: _cell_str(values[i] if i < len(values) else "")
                for i in range(len(unique_headers))
            }
            rows.append(row)
        return unique_headers, rows, sheets
    finally:
        wb.close()


def load_tabular(
    data: bytes,
    filename: str,
    sheet_name: str | None = None,
) -> tuple[str, list[str], list[dict[str, str]], list[str], str | None]:
    fmt = detect_upload_format(filename, data)
    if fmt == "csv":
        headers, rows = _read_csv_rows(data)
        return "csv", headers, rows, [], None
    headers, rows, sheets = _read_xlsx_rows(data, sheet_name)
    selected = sheet_name or (sheets[0] if sheets else None)
    return "xlsx", headers, rows, sheets, selected


def _mapped_values(row: dict[str, str], mapping: dict[str, str | None]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for field_name in CANONICAL_FIELDS:
        header = mapping.get(field_name)
        raw = row.get(header, "") if header else ""
        if field_name == "alternate_codes":
            out[field_name] = ProductCatalog._parse_alternates(raw)
        elif field_name == "category_id":
            text = str(raw or "").strip()
            if not text:
                out[field_name] = None
            else:
                try:
                    out[field_name] = int(float(text)) if re.fullmatch(r"\d+(\.0+)?", text) else None
                    if out[field_name] is None and text:
                        out[field_name] = text  # keep for validation error
                except ValueError:
                    out[field_name] = text
        elif field_name == "primary_code":
            out[field_name] = normalize_code(raw)
        else:
            out[field_name] = str(raw or "").strip()
    return out


def _is_blank(values: dict[str, Any]) -> bool:
    if values.get("primary_code"):
        return False
    if values.get("product_name"):
        return False
    if values.get("manufacturer") or values.get("catalog_number") or values.get("cas_number"):
        return False
    if values.get("alternate_codes"):
        return False
    if values.get("category_id") not in (None, ""):
        return False
    return True


def _exact_duplicate(existing: dict[str, Any], values: dict[str, Any]) -> bool:
    alts_a = existing.get("alternate_codes") or []
    alts_b = values.get("alternate_codes") or []
    return (
        (existing.get("product_name") or "") == (values.get("product_name") or "")
        and (existing.get("manufacturer") or "") == (values.get("manufacturer") or "")
        and (existing.get("catalog_number") or "") == (values.get("catalog_number") or "")
        and (existing.get("cas_number") or "") == (values.get("cas_number") or "")
        and normalize_code(existing.get("primary_code")) == normalize_code(values.get("primary_code"))
        and [normalize_code(x) for x in alts_a] == [normalize_code(x) for x in alts_b]
        and existing.get("category_id") == values.get("category_id")
    )


def _metadata_conflict(existing: dict[str, Any], values: dict[str, Any]) -> list[str]:
    conflicts: list[str] = []
    for key in ("product_name", "manufacturer", "catalog_number", "cas_number", "primary_code"):
        old = existing.get(key) or ""
        new = values.get(key) or ""
        if key == "primary_code":
            old = normalize_code(old)
            new = normalize_code(new)
        if old and new and old != new:
            conflicts.append(key)
    return conflicts


def prepare_import(
    catalog: ProductCatalog,
    data: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    mapping: dict[str, str | None] | None = None,
    known_category_ids: set[int] | None = None,
) -> ImportPreview:
    fmt, headers, raw_rows, sheets, selected = load_tabular(data, filename, sheet_name)
    suggested = suggest_mapping(headers)
    effective = dict(suggested)
    if mapping:
        for key, value in mapping.items():
            if key in CANONICAL_FIELDS:
                effective[key] = value

    preview = ImportPreview(
        filename=filename,
        format=fmt,
        sheets=sheets,
        selected_sheet=selected,
        headers=headers,
        suggested_mapping=suggested,
        mapping=effective,
        total_rows=len(raw_rows),
        blank_rows=0,
        valid_rows=0,
        new_products=0,
        updates=0,
        exact_duplicates=0,
        conflicts=0,
        skipped=0,
        errors=[],
        rows=[],
    )

    # Track codes / identities within the file for duplicate detection
    seen_codes: dict[str, int] = {}
    seen_identity: dict[tuple[str, str], int] = {}

    for idx, raw in enumerate(raw_rows, start=2):  # header is row 1
        values = _mapped_values(raw, effective)
        prepared = PreparedRow(row_number=idx, values=values, action="new")

        if _is_blank(values):
            prepared.action = "skip"
            preview.blank_rows += 1
            preview.skipped += 1
            preview.rows.append(prepared)
            continue

        # Identity requirements: need code OR (manufacturer + catalog)
        code = normalize_code(values.get("primary_code"))
        mfg = (values.get("manufacturer") or "").strip()
        cat = (values.get("catalog_number") or "").strip()
        if not code and not (mfg and cat):
            issue = RowIssue(idx, "identity", "Missing required identity (primary_code or manufacturer+catalog_number)")
            prepared.action = "error"
            prepared.issues.append(issue)
            preview.errors.append(issue)
            preview.rows.append(prepared)
            continue

        cat_id = values.get("category_id")
        if cat_id is not None and not isinstance(cat_id, int):
            issue = RowIssue(idx, "category_id", "Invalid category reference")
            prepared.action = "error"
            prepared.issues.append(issue)
            preview.errors.append(issue)
            preview.rows.append(prepared)
            continue
        if (
            isinstance(cat_id, int)
            and known_category_ids is not None
            and cat_id not in known_category_ids
        ):
            issue = RowIssue(idx, "category_id", f"Unknown category_id {cat_id}")
            prepared.action = "error"
            prepared.issues.append(issue)
            preview.errors.append(issue)
            preview.rows.append(prepared)
            continue

        # Intra-file duplicate codes
        if code:
            if code in seen_codes:
                issue = RowIssue(idx, "primary_code", f"Duplicate code also on row {seen_codes[code]}")
                prepared.action = "conflict"
                prepared.issues.append(issue)
                preview.errors.append(issue)
                preview.conflicts += 1
                preview.rows.append(prepared)
                continue
            seen_codes[code] = idx
            for alt in values.get("alternate_codes") or []:
                alt_n = normalize_code(alt)
                if alt_n and alt_n in seen_codes:
                    issue = RowIssue(idx, "alternate_codes", f"Duplicate code also on row {seen_codes[alt_n]}")
                    prepared.action = "conflict"
                    prepared.issues.append(issue)
                    preview.errors.append(issue)
                    preview.conflicts += 1
                    preview.rows.append(prepared)
                    break
                if alt_n:
                    seen_codes[alt_n] = idx
            if prepared.action == "conflict":
                continue

        identity_key = (normalize_identity_part(mfg), normalize_identity_part(cat))
        if identity_key[0] and identity_key[1]:
            if identity_key in seen_identity:
                issue = RowIssue(
                    idx,
                    "manufacturer+catalog_number",
                    f"Duplicate manufacturer/catalog also on row {seen_identity[identity_key]}",
                )
                prepared.action = "conflict"
                prepared.issues.append(issue)
                preview.errors.append(issue)
                preview.conflicts += 1
                preview.rows.append(prepared)
                continue
            seen_identity[identity_key] = idx

        # Existing catalog matches
        existing_by_code = catalog.find_by_code(code) if code else []
        existing_by_id = (
            catalog.find_by_manufacturer_catalog(mfg, cat) if mfg and cat else []
        )
        existing_ids = {p.id for p in existing_by_code + existing_by_id}
        if len(existing_ids) > 1:
            issue = RowIssue(idx, "identity", "Conflicts with multiple existing catalog products")
            prepared.action = "conflict"
            prepared.issues.append(issue)
            preview.errors.append(issue)
            preview.conflicts += 1
            preview.rows.append(prepared)
            continue

        if len(existing_ids) == 1:
            existing = (existing_by_code or existing_by_id)[0]
            prepared.existing_id = existing.id
            existing_dict = existing.to_dict()
            if _exact_duplicate(existing_dict, values):
                prepared.action = "duplicate"
                preview.exact_duplicates += 1
            else:
                conflicts = _metadata_conflict(existing_dict, values)
                # Also check code collision belonging to a different product
                if code:
                    code_hits = catalog.find_by_code(code)
                    if any(h.id != existing.id for h in code_hits):
                        conflicts.append("primary_code")
                if conflicts:
                    issue = RowIssue(
                        idx,
                        ",".join(conflicts),
                        "Conflicting metadata vs existing catalog product — will not silently merge",
                    )
                    prepared.action = "conflict"
                    prepared.issues.append(issue)
                    preview.errors.append(issue)
                    preview.conflicts += 1
                else:
                    prepared.action = "update"
                    preview.updates += 1
                    preview.valid_rows += 1
            preview.rows.append(prepared)
            continue

        prepared.action = "new"
        preview.new_products += 1
        preview.valid_rows += 1
        preview.rows.append(prepared)

    return preview


def commit_import(
    catalog: ProductCatalog,
    data: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    mapping: dict[str, str | None] | None = None,
    known_category_ids: set[int] | None = None,
    confirm: bool = False,
) -> dict[str, Any]:
    if not confirm:
        raise ValueError("Import commit requires explicit confirmation")

    preview = prepare_import(
        catalog,
        data,
        filename,
        sheet_name=sheet_name,
        mapping=mapping,
        known_category_ids=known_category_ids,
    )
    fatal = [r for r in preview.rows if r.action in {"conflict", "error"}]
    if fatal:
        return {
            "committed": False,
            "reason": "validation_failed",
            "summary": preview.summary(),
            "error_report_csv": preview.error_report_csv(),
        }

    applied_new = 0
    applied_updates = 0
    skipped = 0
    with catalog._conn() as conn:
        try:
            for row in preview.rows:
                if row.action in {"skip", "duplicate"}:
                    skipped += 1
                    continue
                values = row.values
                if row.action == "new":
                    catalog.insert(
                        product_name=values.get("product_name") or "",
                        manufacturer=values.get("manufacturer") or "",
                        catalog_number=values.get("catalog_number") or "",
                        cas_number=values.get("cas_number") or "",
                        primary_code=values.get("primary_code") or "",
                        alternate_codes=values.get("alternate_codes") or [],
                        source=f"import:{filename}",
                        source_row=row.row_number,
                        category_id=values.get("category_id") if isinstance(values.get("category_id"), int) else None,
                        conn=conn,
                    )
                    applied_new += 1
                elif row.action == "update" and row.existing_id:
                    catalog.update(
                        row.existing_id,
                        product_name=values.get("product_name") or "",
                        manufacturer=values.get("manufacturer") or "",
                        catalog_number=values.get("catalog_number") or "",
                        cas_number=values.get("cas_number") or "",
                        primary_code=values.get("primary_code") or "",
                        alternate_codes=values.get("alternate_codes") or [],
                        source=f"import:{filename}",
                        source_row=row.row_number,
                        category_id=values.get("category_id") if isinstance(values.get("category_id"), int) else None,
                        conn=conn,
                    )
                    applied_updates += 1
                else:
                    raise RuntimeError(f"Unexpected row action during commit: {row.action}")
            conn.commit()
        except Exception:
            conn.rollback()
            raise

    return {
        "committed": True,
        "applied_new": applied_new,
        "applied_updates": applied_updates,
        "skipped": skipped,
        "summary": preview.summary(),
        "error_report_csv": preview.error_report_csv(),
    }


def read_upload(file_obj: BinaryIO) -> bytes:
    data = file_obj.read()
    if not data:
        raise ValueError("Empty upload")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"Upload exceeds {MAX_UPLOAD_BYTES} byte limit")
    return data


def mapping_from_form(raw: str | None) -> dict[str, str | None] | None:
    if not raw or not str(raw).strip():
        return None
    parsed = json.loads(raw)
    if not isinstance(parsed, dict):
        raise ValueError("mapping must be a JSON object")
    out: dict[str, str | None] = {}
    for key, value in parsed.items():
        if key in CANONICAL_FIELDS:
            out[key] = None if value in ("", None) else str(value)
    return out
